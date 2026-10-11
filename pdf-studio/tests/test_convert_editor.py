import base64
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover
    import fitz

from app import convert, editor

RO = "Știință și țară: ăâîșț ĂÂÎȘȚ"


def make_pdf(pages=3, text=True):
    doc = fitz.open()
    for i in range(pages):
        p = doc.new_page(width=595, height=842)
        if text:
            p.insert_text((72, 100), f"Pagina {i + 1} Hello World", fontsize=14)
            p.insert_text((72, 200), "SECRET-DATA 12345", fontsize=12)
    data = doc.tobytes()
    doc.close()
    return data


def open_pdf(b):
    return fitz.open(stream=b, filetype="pdf")


def page_text(b, i=0):
    with open_pdf(b) as d:
        return d[i].get_text()


def png_b64():
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 10), False)
    pix.set_rect(pix.irect, (255, 0, 0))
    return base64.b64encode(pix.tobytes("png")).decode()


# ---------------------------------------------------------------- info / render

def test_info():
    d = editor.info(make_pdf(2))
    assert d["pages"] == 2
    assert d["page_sizes"] == [[595, 842], [595, 842]]
    assert d["encrypted"] is False
    assert d["signature_count"] == 0
    assert isinstance(d["metadata"], dict)


def test_info_invalid():
    with pytest.raises(ValueError):
        editor.info(b"not a pdf")


def test_render_page():
    png = editor.render_page(make_pdf(2), 1, zoom=1.0)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    pix = fitz.Pixmap(png)
    assert (pix.width, pix.height) == (595, 842)
    with pytest.raises(ValueError):
        editor.render_page(make_pdf(1), 5)


# ---------------------------------------------------------------- ops

def test_add_text_diacritics():
    out = editor.apply_edits(make_pdf(1), [
        {"type": "add_text", "page": 0, "x": 100, "y": 300, "text": RO, "size": 16, "color": "#ff0000"}])
    with open_pdf(out) as d:
        txt = d[0].get_text()
        assert RO in txt
        blocks = d[0].get_text("dict")["blocks"]
        spans = [s for b in blocks for l in b.get("lines", []) for s in l["spans"] if "Știință" in s["text"]]
        assert spans and spans[0]["color"] == 0xFF0000
        assert abs(spans[0]["size"] - 16) < 0.5
        # top-left origin: text box starts near y=300
        assert 295 <= spans[0]["bbox"][1] <= 305
        assert 99 <= spans[0]["bbox"][0] <= 101


def test_add_text_on_rotated_page():
    out = editor.apply_edits(make_pdf(1), [
        {"type": "rotate", "page": 0, "angle": 90},
        {"type": "add_text", "page": 0, "x": 50, "y": 50, "text": "Rotit ș", "size": 12}])
    assert "Rotit ș" in page_text(out)


def test_highlight():
    out = editor.apply_edits(make_pdf(1), [{"type": "highlight", "page": 0, "rect": [70, 85, 250, 105]}])
    with open_pdf(out) as d:
        pg = d[0]
        annots = list(pg.annots())
        assert len(annots) == 1 and annots[0].type[1] == "Highlight"


def test_rect_and_whiteout():
    out = editor.apply_edits(make_pdf(1), [
        {"type": "rect", "page": 0, "rect": [10, 10, 100, 100], "color": "#0000ff", "fill": "#00ff00", "width": 2},
        {"type": "whiteout", "page": 0, "rect": [300, 300, 400, 400]}])
    with open_pdf(out) as d:
        drawings = d[0].get_drawings()
        fills = [dr.get("fill") for dr in drawings]
        assert any(f and abs(f[1] - 1) < 0.01 and f[0] < 0.01 for f in fills)
        assert any(f and all(abs(c - 1) < 0.01 for c in f) for f in fills)


def test_redact_removes_content():
    pdf = make_pdf(1)
    assert "SECRET-DATA" in page_text(pdf)
    out = editor.apply_edits(pdf, [{"type": "redact", "page": 0, "rect": [60, 180, 400, 210]}])
    txt = page_text(out)
    assert "SECRET" not in txt
    assert "Hello World" in txt
    assert b"SECRET" not in fitz.open(stream=out).tobytes(expand=255)


def test_image_plain_and_data_uri():
    b64 = png_b64()
    out = editor.apply_edits(make_pdf(1), [
        {"type": "image", "page": 0, "rect": [100, 400, 300, 500], "image_b64": b64},
        {"type": "image", "page": 0, "rect": [100, 600, 300, 700], "image_b64": "data:image/png;base64," + b64}])
    with open_pdf(out) as d:
        assert len(d[0].get_images()) >= 1
        infos = d[0].get_image_info()
        assert len(infos) == 2
    with pytest.raises(ValueError):
        editor.apply_edits(make_pdf(1), [{"type": "image", "page": 0, "rect": [0, 0, 10, 10], "image_b64": ""}])


def test_replace_text():
    out = editor.apply_edits(make_pdf(2), [{"type": "replace_text", "search": "Hello World", "replace": "Bună ziua ș"}])
    for i in range(2):
        t = page_text(out, i)
        assert "Hello World" not in t
        assert "Bună ziua ș" in t
        assert "SECRET-DATA" in t
    out2 = editor.apply_edits(make_pdf(2), [{"type": "replace_text", "search": "Hello", "replace": "X", "page": 1}])
    assert "Hello" in page_text(out2, 0) and "Hello" not in page_text(out2, 1)
    with pytest.raises(ValueError):
        editor.apply_edits(make_pdf(1), [{"type": "replace_text", "search": "nuexista", "replace": "x"}])


# ---------------------------------------------------------------- spans / edit_text

def _pdf_with_text(text, x=50, y=100, size=16, w=400, hgt=300):
    doc = fitz.open()
    p = doc.new_page(width=w, height=hgt)
    p.insert_text((x, y), text, fontsize=size)
    data = doc.tobytes()
    doc.close()
    return data


def test_spans_basic():
    res = editor.spans(_pdf_with_text("Vechi text", x=50, y=100, size=14, w=300, hgt=200), 0)
    assert res["width"] == 300 and res["height"] == 200
    found = [s for s in res["spans"] if "Vechi" in s["text"]]
    assert found, "span-ul cu textul cunoscut nu a fost găsit"
    s = found[0]
    assert abs(s["size"] - 14) < 1.0
    x0, y0, x1, y1 = s["bbox"]
    # bbox plauzibil: originea sus-stânga, lângă punctul de inserare (x~50, baseline~100)
    assert 0 <= x0 <= 300 and 0 <= y0 <= 200 and x1 <= 300 and y1 <= 200
    assert 40 <= x0 <= 60
    assert y0 <= 100 <= y1
    assert s["color"].startswith("#") and len(s["color"]) == 7
    assert isinstance(s["bold"], bool) and isinstance(s["italic"], bool)


def test_spans_skips_whitespace():
    doc = fitz.open()
    p = doc.new_page(width=300, height=200)
    p.insert_text((50, 50), "   ", fontsize=12)
    p.insert_text((50, 120), "Real", fontsize=12)
    data = doc.tobytes(); doc.close()
    res = editor.spans(data, 0)
    assert res["spans"]
    assert all(s["text"].strip() for s in res["spans"])


def test_spans_bad_page():
    with pytest.raises(ValueError):
        editor.spans(_pdf_with_text("x"), 9)


def test_edit_text_replaces():
    data = _pdf_with_text("Vechi")
    sp = [s for s in editor.spans(data, 0)["spans"] if "Vechi" in s["text"]][0]
    out = editor.apply_edits(data, [{
        "type": "edit_text", "page": 0, "rect": sp["bbox"],
        "text": "Nou", "size": sp["size"], "color": sp["color"],
    }])
    t = page_text(out, 0)
    assert "Nou" in t
    assert "Vechi" not in t


def test_edit_text_diacritics():
    data = _pdf_with_text("Vechi")
    sp = [s for s in editor.spans(data, 0)["spans"] if "Vechi" in s["text"]][0]
    out = editor.apply_edits(data, [{
        "type": "edit_text", "page": 0, "rect": sp["bbox"],
        "text": "Țară șiț", "size": sp["size"], "color": "#ff0000",
    }])
    t = page_text(out, 0)
    assert "Țară" in t and "șiț" in t
    assert "Vechi" not in t


def test_edit_text_coexists_with_other_ops():
    data = _pdf_with_text("Vechi")
    sp = [s for s in editor.spans(data, 0)["spans"] if "Vechi" in s["text"]][0]
    out = editor.apply_edits(data, [
        {"type": "edit_text", "page": 0, "rect": sp["bbox"], "text": "Nou", "size": sp["size"], "color": sp["color"]},
        {"type": "add_text", "page": 0, "x": 50, "y": 200, "text": "Adăugat ș", "size": 12, "color": "#0000ff"},
        {"type": "highlight", "page": 0, "rect": [40, 180, 200, 215]},
    ])
    t = page_text(out, 0)
    assert "Nou" in t and "Adăugat ș" in t and "Vechi" not in t


def test_rotate():
    out = editor.apply_edits(make_pdf(2), [{"type": "rotate", "page": 1, "angle": 90}])
    with open_pdf(out) as d:
        assert d[0].rotation == 0 and d[1].rotation == 90
    with pytest.raises(ValueError):
        editor.apply_edits(make_pdf(1), [{"type": "rotate", "page": 0, "angle": 45}])


def test_delete_page():
    out = editor.apply_edits(make_pdf(3), [{"type": "delete_page", "page": 1}])
    assert editor.info(out)["pages"] == 2
    assert "Pagina 3" in page_text(out, 1)


@pytest.mark.parametrize("frm,to,expected", [
    (0, 2, ["2", "3", "1"]),
    (2, 0, ["3", "1", "2"]),
    (0, 1, ["2", "1", "3"]),
    (1, 2, ["1", "3", "2"]),
])
def test_move_page(frm, to, expected):
    out = editor.apply_edits(make_pdf(3), [{"type": "move_page", "from": frm, "to": to}])
    got = [page_text(out, i).split("Pagina ")[1][0] for i in range(3)]
    assert got == expected


def test_insert_blank():
    out = editor.apply_edits(make_pdf(2), [{"type": "insert_blank", "at": 1}, {"type": "insert_blank", "at": 3}])
    assert editor.info(out)["pages"] == 4
    assert page_text(out, 1).strip() == "" and page_text(out, 3).strip() == ""
    assert "Pagina 2" in page_text(out, 2)


def test_note():
    out = editor.apply_edits(make_pdf(1), [{"type": "note", "page": 0, "x": 50, "y": 50, "text": "Notă: verificați"}])
    with open_pdf(out) as d:
        pg = d[0]
        a = list(pg.annots())
        assert a[0].type[1] == "Text" and a[0].info["content"] == "Notă: verificați"


def test_set_metadata():
    out = editor.apply_edits(make_pdf(1), [{"type": "set_metadata", "title": "Contract ș", "author": "Ion Țăranu",
                                            "subject": "Test", "keywords": "a, b"}])
    m = editor.info(out)["metadata"]
    assert m["title"] == "Contract ș" and m["author"] == "Ion Țăranu"
    assert m["subject"] == "Test" and m["keywords"] == "a, b"


def test_unknown_op_and_bad_page():
    with pytest.raises(ValueError):
        editor.apply_edits(make_pdf(1), [{"type": "zbor"}])
    with pytest.raises(ValueError):
        editor.apply_edits(make_pdf(1), [{"type": "add_text", "page": 3, "x": 1, "y": 1, "text": "x"}])


# ---------------------------------------------------------------- merge / split / text / compress

def test_merge():
    out = editor.merge([make_pdf(2), make_pdf(3)])
    assert editor.info(out)["pages"] == 5
    with pytest.raises(ValueError):
        editor.merge([])


def test_split():
    parts = editor.split(make_pdf(6), "1-3, 4,5-end")
    assert [editor.info(p)["pages"] for p in parts] == [3, 1, 2]
    assert "Pagina 4" in page_text(parts[1])
    assert "Pagina 5" in page_text(parts[2])
    for bad in ["0-2", "7", "3-1", "abc", ""]:
        with pytest.raises(ValueError):
            editor.split(make_pdf(6), bad)


def test_extract_text():
    t = editor.extract_text(make_pdf(2))
    assert "Pagina 1" in t and "Pagina 2 Hello World" in t


def test_compress():
    doc = fitz.open()
    for _ in range(5):
        p = doc.new_page()
        p.insert_text((72, 72), "Compresie " * 50)
    raw = doc.tobytes(garbage=0, deflate=False)
    out = editor.compress(raw)
    assert len(out) < len(raw)
    assert editor.info(out)["pages"] == 5


# ---------------------------------------------------------------- convert

def test_find_soffice():
    p = convert.find_soffice()
    assert p is None or os.path.exists(p)


def test_office_to_pdf_rejects_unknown():
    with pytest.raises(ValueError):
        convert.office_to_pdf(b"abc", "file.exe")
    with pytest.raises(ValueError):
        convert.office_to_pdf(b"", "file.docx")


def test_pdf_to_docx():
    out = convert.pdf_to_docx(make_pdf(1))
    assert out[:2] == b"PK"
    import docx
    d = docx.Document(io.BytesIO(out))
    assert "Hello World" in "\n".join(p.text for p in d.paragraphs)


@pytest.mark.skipif(convert.find_soffice() is None, reason="LibreOffice indisponibil")
def test_round_trip_docx_pdf_docx():
    import docx
    d = docx.Document()
    d.add_heading("Contract de prestări servicii", 1)
    d.add_paragraph(RO)
    buf = io.BytesIO()
    d.save(buf)
    pdf = convert.office_to_pdf(buf.getvalue(), "contract ș.docx")
    assert pdf[:5] == b"%PDF-"
    txt = editor.extract_text(pdf)
    assert "Știință și țară" in txt and "ĂÂÎȘȚ" in txt
    back = convert.pdf_to_docx(pdf)
    d2 = docx.Document(io.BytesIO(back))
    all_text = "\n".join(p.text for p in d2.paragraphs)
    assert "Știință" in all_text and "ăâîșț" in all_text
    pdf2 = convert.office_to_pdf(back, "back.docx")
    assert "Știință" in editor.extract_text(pdf2)


@pytest.mark.skipif(convert.find_soffice() is None, reason="LibreOffice indisponibil")
def test_office_to_pdf_txt():
    pdf = convert.office_to_pdf("Salut ăâîșț\n".encode("utf-8"), "nota.txt")
    assert "Salut" in editor.extract_text(pdf)
